"""Importação de despesas (CSV/XLSX) e cadastro manual (RF01 / US01).

Regras:

- o cabeçalho é obrigatório e aceita variações de acento, maiúsculas e espaços
  ("Conta Contábil" → ``conta_contabil``). Colunas extras são ignoradas;
- se faltar uma coluna obrigatória, o arquivo inteiro é recusado
  (``ArquivoInvalidoError``), sem criar lote;
- cada linha é validada à parte. As válidas são importadas e as inválidas ficam
  em ``LoteImportacao.erros`` como ``{"linha", "campo", "mensagem"}``. A
  numeração segue a do Excel: o cabeçalho é a linha 1;
- CSV com ``,`` ou ``;`` como separador. Valores em ``1234.56``, ``1234,56`` ou
  ``1.234,56``. Datas em ``AAAA-MM-DD`` ou ``DD/MM/AAAA`` (ou célula de data no XLSX);
- despesas repetidas não são bloqueadas: duplicidade é uma das anomalias que o
  motor deve detectar.

As funções não fazem commit; quem chama (rota ou comando) decide.
"""

import csv
import io
import re
import unicodedata
import zipfile
from datetime import date, datetime
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import PurePath

from openpyxl import load_workbook
from sqlalchemy import insert

from app.extensoes import db
from app.models import Despesa, LoteImportacao, Usuario
from app.servicos.estatisticas import recalcular_estatisticas

CAMPOS_OBRIGATORIOS = (
    "valor",
    "data",
    "categoria",
    "conta_contabil",
    "centro_custo",
    "funcionario",
)
CAMPOS = CAMPOS_OBRIGATORIOS + ("descricao",)
CAMPOS_TEXTO = ("categoria", "conta_contabil", "centro_custo", "funcionario")
EXTENSOES = (".csv", ".xlsx")

NOMES_CAMPOS = {
    "valor": "valor",
    "data": "data",
    "categoria": "categoria",
    "conta_contabil": "conta contábil",
    "centro_custo": "centro de custo",
    "funcionario": "funcionário",
    "descricao": "descrição",
}

# Variações de cabeçalho aceitas, já normalizadas (sem acento, minúsculas, "_").
APELIDOS_CABECALHO = {
    "valor_r": "valor",
    "conta": "conta_contabil",
    "centro_de_custo": "centro_custo",
    "cc": "centro_custo",
}

# Limite de erros guardados por lote, para não inflar o JSON com arquivos ruins.
MAXIMO_ERROS_REGISTRADOS = 500

# Numeric(14, 2): até 12 dígitos antes da vírgula.
_VALOR_MAXIMO = Decimal("1e12")
_CENTAVOS = Decimal("0.01")
_FORMATOS_DATA = ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%Y-%m-%d %H:%M:%S")
_MILHAR_COM_PONTO = re.compile(r"^\d{1,3}(\.\d{3})+$")


class ArquivoInvalidoError(ValueError):
    """O arquivo não pode ser importado (formato, cabeçalho ou conteúdo)."""


class DespesaInvalidaError(ValueError):
    """Dados de uma despesa inválidos. ``erros`` é uma lista de (campo, mensagem)."""

    def __init__(self, erros: list[tuple[str, str]]):
        super().__init__("; ".join(mensagem for _, mensagem in erros))
        self.erros = erros


# --- Validação de uma despesa -------------------------------------------------


def _tamanho_maximo(campo: str) -> int | None:
    return getattr(Despesa.__table__.c[campo].type, "length", None)


def _vazio(valor) -> bool:
    return valor is None or (isinstance(valor, str) and not valor.strip())


def _texto(valor) -> str:
    if isinstance(valor, float) and valor.is_integer():
        valor = int(valor)  # ex.: conta 3101 lida do Excel como 3101.0
    return str(valor).strip()


def converter_valor(bruto, decimal_com_virgula: bool = False) -> Decimal:
    """Converte o valor para Decimal com centavos. Levanta ValueError se inválido.

    ``decimal_com_virgula`` indica um arquivo no formato brasileiro (CSV com
    ``;``), em que ``1.234`` significa mil duzentos e trinta e quatro.
    """
    if isinstance(bruto, bool):
        raise ValueError("valor inválido")
    if isinstance(bruto, int | float | Decimal):
        texto = str(bruto)
    else:
        texto = re.sub(r"\s", "", str(bruto)).replace("R$", "")
        if "," in texto and "." in texto:
            # O último separador é o decimal; o outro é o de milhar.
            milhar = "." if texto.rfind(",") > texto.rfind(".") else ","
            texto = texto.replace(milhar, "")
        elif "." in texto and decimal_com_virgula and _MILHAR_COM_PONTO.match(texto):
            texto = texto.replace(".", "")
        texto = texto.replace(",", ".")
    try:
        valor = Decimal(texto)
    except InvalidOperation as erro:
        raise ValueError(f"valor inválido: {str(bruto).strip()!r}") from erro
    if not valor.is_finite():
        raise ValueError(f"valor inválido: {str(bruto).strip()!r}")
    valor = valor.quantize(_CENTAVOS, ROUND_HALF_UP)
    if valor <= 0:
        raise ValueError("o valor deve ser maior que zero")
    if valor >= _VALOR_MAXIMO:
        raise ValueError("valor alto demais")
    return valor


def converter_data(bruto, hoje: date) -> date:
    """Converte a data. Levanta ValueError se inválida ou futura."""
    if isinstance(bruto, datetime):
        dia = bruto.date()
    elif isinstance(bruto, date):
        dia = bruto
    else:
        texto = str(bruto).strip()
        for formato in _FORMATOS_DATA:
            try:
                dia = datetime.strptime(texto, formato).date()
                break
            except ValueError:
                continue
        else:
            raise ValueError(f"data inválida: {texto!r} (use AAAA-MM-DD ou DD/MM/AAAA)")
    if dia > hoje:
        raise ValueError(f"data futura: {dia.strftime('%d/%m/%Y')}")
    return dia


def validar_despesa(
    bruto: dict, hoje: date | None = None, decimal_com_virgula: bool = False
) -> tuple[dict | None, list[tuple[str, str]]]:
    """Valida os dados brutos de uma despesa.

    Retorna ``(dados, [])`` com os valores convertidos, ou ``(None, erros)``
    com a lista de (campo, mensagem).
    """
    hoje = hoje or date.today()
    dados, erros = {}, []
    for campo in CAMPOS_OBRIGATORIOS:
        valor = bruto.get(campo)
        if _vazio(valor):
            erros.append((campo, f"{NOMES_CAMPOS[campo]} é obrigatório"))
            continue
        try:
            if campo == "valor":
                dados[campo] = converter_valor(valor, decimal_com_virgula)
            elif campo == "data":
                dados[campo] = converter_data(valor, hoje)
            else:
                texto = _texto(valor)
                maximo = _tamanho_maximo(campo)
                if maximo and len(texto) > maximo:
                    raise ValueError(f"{NOMES_CAMPOS[campo]} tem mais de {maximo} caracteres")
                dados[campo] = texto
        except ValueError as erro:
            erros.append((campo, str(erro)))
    descricao = bruto.get("descricao")
    dados["descricao"] = None if _vazio(descricao) else _texto(descricao)
    return (None, erros) if erros else (dados, [])


# --- Leitura dos arquivos -----------------------------------------------------


def normalizar_cabecalho(nome) -> str:
    """Normaliza um nome de coluna ("Conta Contábil" → "conta_contabil") e aplica os apelidos."""
    texto = unicodedata.normalize("NFKD", str(nome or "")).encode("ascii", "ignore").decode()
    texto = re.sub(r"[^a-z0-9]+", "_", texto.lower()).strip("_")
    return APELIDOS_CABECALHO.get(texto, texto)


def _mapear_cabecalho(cabecalho: list) -> dict[str, int]:
    """Campo → índice da coluna. Recusa colunas obrigatórias ausentes ou repetidas."""
    mapa: dict[str, int] = {}
    for indice, nome in enumerate(cabecalho):
        campo = normalizar_cabecalho(nome)
        if campo not in CAMPOS:
            continue
        if campo in mapa:
            raise ArquivoInvalidoError(f"A coluna {NOMES_CAMPOS[campo]} aparece mais de uma vez.")
        mapa[campo] = indice
    faltando = [NOMES_CAMPOS[c] for c in CAMPOS_OBRIGATORIOS if c not in mapa]
    if faltando:
        raise ArquivoInvalidoError(
            "Colunas obrigatórias ausentes: " + ", ".join(faltando) + ". "
            "O arquivo precisa de: valor, data, categoria, conta contábil, "
            "centro de custo e funcionário."
        )
    return mapa


def _decodificar(conteudo: bytes) -> str:
    for codificacao in ("utf-8-sig", "cp1252"):  # cp1252: CSV salvo pelo Excel no Windows
        try:
            return conteudo.decode(codificacao)
        except UnicodeDecodeError:
            continue
    raise ArquivoInvalidoError("Não foi possível ler o texto do arquivo (use UTF-8).")


def _ler_csv(conteudo: bytes) -> tuple[list[list], bool]:
    """Retorna as linhas e se o arquivo usa o formato brasileiro (``;`` e vírgula decimal)."""
    texto = _decodificar(conteudo)
    primeira_linha = texto.split("\n", 1)[0]
    separador = max((";", ",", "\t"), key=primeira_linha.count)
    linhas = list(csv.reader(io.StringIO(texto, newline=""), delimiter=separador))
    return linhas, separador == ";"


def _ler_xlsx(conteudo: bytes) -> list[list]:
    try:
        planilha = load_workbook(io.BytesIO(conteudo), read_only=True, data_only=True)
    except (zipfile.BadZipFile, KeyError, OSError, ValueError) as erro:
        raise ArquivoInvalidoError("O arquivo não é uma planilha XLSX válida.") from erro
    try:
        return [list(linha) for linha in planilha.worksheets[0].iter_rows(values_only=True)]
    finally:
        planilha.close()


def ler_arquivo(nome_arquivo: str, conteudo: bytes) -> tuple[dict[str, int], list, bool]:
    """Lê o arquivo e devolve (mapa do cabeçalho, linhas de dados, formato brasileiro)."""
    extensao = PurePath(nome_arquivo).suffix.lower()
    if extensao not in EXTENSOES:
        raise ArquivoInvalidoError("Formato não suportado. Envie um arquivo .csv ou .xlsx.")
    if extensao == ".csv":
        linhas, formato_brasileiro = _ler_csv(conteudo)
    else:
        linhas, formato_brasileiro = _ler_xlsx(conteudo), False
    if not linhas or all(_vazio(v) for v in linhas[0]):
        raise ArquivoInvalidoError("O arquivo está vazio ou não tem cabeçalho na primeira linha.")
    return _mapear_cabecalho(linhas[0]), linhas[1:], formato_brasileiro


# --- Casos de uso -------------------------------------------------------------


def importar_arquivo(
    nome_arquivo: str, conteudo: bytes, usuario: Usuario, hoje: date | None = None
) -> LoteImportacao:
    """Importa as linhas válidas num novo lote e recalcula as estatísticas (US02).

    Levanta ``ArquivoInvalidoError`` se o arquivo não puder ser lido.
    """
    hoje = hoje or date.today()
    mapa, linhas, formato_brasileiro = ler_arquivo(nome_arquivo, conteudo)

    despesas, erros, total = [], [], 0
    for numero, linha in enumerate(linhas, start=2):
        bruto = {campo: linha[i] if i < len(linha) else None for campo, i in mapa.items()}
        if all(_vazio(v) for v in bruto.values()):
            continue  # linha em branco
        total += 1
        dados, erros_linha = validar_despesa(bruto, hoje, formato_brasileiro)
        if dados:
            despesas.append(dados)
        erros.extend({"linha": numero, "campo": c, "mensagem": m} for c, m in erros_linha)

    if len(erros) > MAXIMO_ERROS_REGISTRADOS:
        omitidos = len(erros) - MAXIMO_ERROS_REGISTRADOS
        erros = erros[:MAXIMO_ERROS_REGISTRADOS] + [
            {"linha": None, "campo": None, "mensagem": f"... e mais {omitidos} erros."}
        ]

    lote = LoteImportacao(
        nome_arquivo=PurePath(nome_arquivo).name[:255],
        importado_por=usuario.id,
        total_linhas=total,
        linhas_validas=len(despesas),
        erros=erros,
    )
    db.session.add(lote)
    db.session.flush()
    if despesas:
        db.session.execute(insert(Despesa), [{**d, "lote_id": lote.id} for d in despesas])
        recalcular_estatisticas()
    return lote


def cadastrar_despesa(bruto: dict, hoje: date | None = None) -> Despesa:
    """Cadastra uma despesa sem lote e recalcula as estatísticas.

    Levanta ``DespesaInvalidaError`` com os erros por campo.
    """
    dados, erros = validar_despesa(bruto, hoje)
    if erros:
        raise DespesaInvalidaError(erros)
    despesa = Despesa(**dados)
    db.session.add(despesa)
    db.session.flush()
    recalcular_estatisticas()
    return despesa
